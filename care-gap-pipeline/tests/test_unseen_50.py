"""
Formal Unseen-Data Test for Medicare Advantage Care-Gap Pipeline.
Tests the generalization of the ML + Intervention Selection + Optimizer pipeline
on data/test/Unseen_50_Test_Data_New_Combinations (1).xlsx without modifying production code.
Generates outputs/unseen_new_combinations_50_test_report.xlsx with 10 detailed validation sheets.
"""

import os
import sys
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# Ensure root directory is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from pipeline.intervention_candidates import generate_intervention_candidates, ACTIONABLE_CHANNELS
from ml_model.infer import run_inference
from pipeline.calibration import calibrate_channel_bias
from pipeline.intervention_selection import select_best_interventions
from pipeline.adapters import ml_to_optimizer_input
from optimizer import run_optimizer, MAX_CHANNEL_SHARE, SELECTION_MARGIN


def run_unseen_50_test(
    test_excel_path: str = "data/test/Unseen_50_Test_Data_New_Combinations (1).xlsx",
    output_report_path: str = "outputs/unseen_new_combinations_50_test_report.xlsx"
):
    print("=" * 60)
    print("STARTING UNSEEN DATASET TEST (50 COMBINATIONS)")
    print("=" * 60)
    print(f"Test File: {test_excel_path}")

    failed_checks_list = []
    
    def record_check(check_id, stage, description, expected, actual, passed, remarks=""):
        res = {
            "Check ID": check_id,
            "Stage": stage,
            "Description": description,
            "Expected": str(expected),
            "Actual": str(actual),
            "Status": "PASS" if passed else "FAIL",
            "Remarks": remarks
        }
        if not passed:
            failed_checks_list.append(res)
        return res

    all_checks = []

    # =========================================================================
    # STEP 2 — VALIDATE THE INPUT DATASET
    # =========================================================================
    print("\n--- STEP 2: VALIDATING INPUT DATASET ---")
    if not os.path.exists(test_excel_path):
        raise FileNotFoundError(f"Test dataset not found at {test_excel_path}")

    raw_input_df = pd.read_excel(test_excel_path, sheet_name="Unseen_New_50")
    num_rows = len(raw_input_df)
    num_cols = len(raw_input_df.columns)
    has_outcome = "outcome" in [c.lower() for c in raw_input_df.columns]
    
    required_id_cols = ["patient_id", "plan_id", "care_gap"]
    missing_req = [c for c in required_id_cols if c not in raw_input_df.columns]
    
    c1 = record_check("CHK_INP_01", "Input Validation", "Row Count == 50", 50, num_rows, num_rows == 50)
    c2 = record_check("CHK_INP_02", "Input Validation", "No Target/Outcome column", False, has_outcome, not has_outcome)
    c3 = record_check("CHK_INP_03", "Input Validation", "Required ID columns present", [], missing_req, len(missing_req) == 0)
    all_checks.extend([c1, c2, c3])

    print("INPUT DATASET VALIDATION")
    print("------------------------")
    print(f"Rows: {num_rows}")
    print(f"Columns: {num_cols}")
    print(f"Outcome column present: {'YES' if has_outcome else 'NO'}")
    print(f"Missing required columns: {missing_req if missing_req else 'None'}")
    input_status = "PASS" if (num_rows == 50 and not has_outcome and len(missing_req) == 0) else "FAIL"
    print(f"Status: {input_status}")

    # =========================================================================
    # STEP 3 — GENERATE INTERVENTION CANDIDATES
    # =========================================================================
    print("\n--- STEP 3: GENERATING INTERVENTION CANDIDATES ---")
    candidate_df = generate_intervention_candidates(raw_input_df)
    total_cands = len(candidate_df)
    ch_counts = candidate_df["intervention_type"].value_counts().to_dict()
    email_cands = ch_counts.get("Email", 0)
    sms_cands = ch_counts.get("SMS", 0)
    phone_cands = ch_counts.get("Phone Call", 0)

    # Check 3 candidates per patient/care_gap
    cands_per_gap = candidate_df.groupby(["patient_id", "care_gap"])["intervention_type"].nunique()
    all_three = (cands_per_gap == 3).all()

    c4 = record_check("CHK_CND_01", "Candidate Gen", "Total Candidates == 150", 150, total_cands, total_cands == 150)
    c5 = record_check("CHK_CND_02", "Candidate Gen", "Each channel count == 50", "Email:50, SMS:50, Phone:50", f"Email:{email_cands}, SMS:{sms_cands}, Phone:{phone_cands}", email_cands == 50 and sms_cands == 50 and phone_cands == 50)
    c6 = record_check("CHK_CND_03", "Candidate Gen", "Exactly 3 distinct channels per context", True, bool(all_three), all_three)
    all_checks.extend([c4, c5, c6])

    print("CANDIDATE GENERATION")
    print("--------------------")
    print(f"Original contexts: {num_rows}")
    print(f"Candidates generated: {total_cands}")
    print(f"Email: {email_cands}")
    print(f"SMS: {sms_cands}")
    print(f"Phone Call: {phone_cands}")
    cand_status = "PASS" if (total_cands == 150 and email_cands == 50 and sms_cands == 50 and phone_cands == 50 and all_three) else "FAIL"
    print(f"Status: {cand_status}")

    # =========================================================================
    # STEP 4 — EXISTING ML INFERENCE
    # =========================================================================
    print("\n--- STEP 4: RUNNING ML INFERENCE ---")
    ml_output_df = run_inference(candidate_df, models_dir="models")
    preds_count = len(ml_output_df)
    missing_probs = ml_output_df["probability_score"].isna().sum()
    invalid_probs = ((ml_output_df["probability_score"] < 0.0) | (ml_output_df["probability_score"] > 1.0)).sum()
    valid_bounds = (ml_output_df["uncertainity_range_lower"] <= ml_output_df["uncertainity_range_upper"]).all()

    c7 = record_check("CHK_ML_01", "ML Inference", "Prediction rows == 150", 150, preds_count, preds_count == 150)
    c8 = record_check("CHK_ML_02", "ML Inference", "Missing probability scores == 0", 0, missing_probs, missing_probs == 0)
    c9 = record_check("CHK_ML_03", "ML Inference", "Invalid probability scores [0, 1] == 0", 0, invalid_probs, invalid_probs == 0)
    c10 = record_check("CHK_ML_04", "ML Inference", "Valid uncertainty bounds (lower <= upper)", True, bool(valid_bounds), valid_bounds)
    all_checks.extend([c7, c8, c9, c10])

    print("ML INFERENCE")
    print("------------")
    print(f"Input candidates: {total_cands}")
    print(f"Predictions: {preds_count}")
    print(f"Missing probabilities: {missing_probs}")
    print(f"Invalid probabilities: {invalid_probs}")
    ml_status = "PASS" if (preds_count == 150 and missing_probs == 0 and invalid_probs == 0 and valid_bounds) else "FAIL"
    print(f"Status: {ml_status}")

    # =========================================================================
    # STEP 5 — LOGIT CALIBRATION
    # =========================================================================
    print("\n--- STEP 5: RUNNING LOGIT CALIBRATION ---")
    calibrated_probs = calibrate_channel_bias(ml_output_df)
    ml_output_df["calibrated_probability"] = calibrated_probs
    candidate_df["calibrated_probability"] = calibrated_probs

    cal_count = len(ml_output_df)
    missing_cal = ml_output_df["calibrated_probability"].isna().sum()
    invalid_cal = ((ml_output_df["calibrated_probability"] < 0.0) | (ml_output_df["calibrated_probability"] > 1.0)).sum()

    c11 = record_check("CHK_CAL_01", "Calibration", "Calibrated rows == 150", 150, cal_count, cal_count == 150)
    c12 = record_check("CHK_CAL_02", "Calibration", "Missing calibrated probabilities == 0", 0, missing_cal, missing_cal == 0)
    c13 = record_check("CHK_CAL_03", "Calibration", "Invalid calibrated probabilities [0, 1] == 0", 0, invalid_cal, invalid_cal == 0)
    all_checks.extend([c11, c12, c13])

    print("CALIBRATION")
    print("-----------")
    print(f"Rows calibrated: {cal_count}")
    print(f"Missing calibrated probabilities: {missing_cal}")
    print(f"Invalid calibrated probabilities: {invalid_cal}")
    cal_status = "PASS" if (cal_count == 150 and missing_cal == 0 and invalid_cal == 0) else "FAIL"
    print(f"Status: {cal_status}")

    # =========================================================================
    # STEP 6 — BEST INTERVENTION SELECTION
    # =========================================================================
    print("\n--- STEP 6: RUNNING BEST INTERVENTION SELECTION ---")
    best_cands_df, audit_df, metrics = select_best_interventions(ml_output_df)
    num_selected = len(best_cands_df)
    clear_winners = metrics.get("clear_winner_decisions", 0)
    near_ties = metrics.get("near_tie_decisions", 0)
    clear_changed = metrics.get("clear_winner_changed_count", 0)
    
    valid_channels = {"Email", "SMS", "Phone Call"}
    actual_sel_channels = set(best_cands_df["intervention_type"].unique())
    invalid_interventions = len(actual_sel_channels - valid_channels)

    c14 = record_check("CHK_SEL_01", "Intervention Selection", "Selected count == 50", 50, num_selected, num_selected == 50)
    c15 = record_check("CHK_SEL_02", "Intervention Selection", "Clear winners changed == 0", 0, clear_changed, clear_changed == 0)
    c16 = record_check("CHK_SEL_03", "Intervention Selection", "Invalid intervention count == 0", 0, invalid_interventions, invalid_interventions == 0)
    all_checks.extend([c14, c15, c16])

    print("BEST INTERVENTION SELECTION")
    print("---------------------------")
    print(f"Input candidates: {len(ml_output_df)}")
    print(f"Selected interventions: {num_selected}")
    print(f"Clear winners: {clear_winners}")
    print(f"Near ties: {near_ties}")
    print(f"Clear winners changed: {clear_changed}")
    print(f"Invalid interventions: {invalid_interventions}")
    sel_status = "PASS" if (num_selected == 50 and clear_changed == 0 and invalid_interventions == 0) else "FAIL"
    print(f"Status: {sel_status}")

    # =========================================================================
    # STEP 7 — CHANNEL DISTRIBUTION
    # =========================================================================
    print("\n--- STEP 7: ANALYZING CHANNEL DISTRIBUTION ---")
    sel_counts = best_cands_df["intervention_type"].value_counts().to_dict()
    email_sel = sel_counts.get("Email", 0)
    sms_sel = sel_counts.get("SMS", 0)
    phone_sel = sel_counts.get("Phone Call", 0)
    
    email_pct = (email_sel / num_selected) * 100.0 if num_selected > 0 else 0.0
    sms_pct = (sms_sel / num_selected) * 100.0 if num_selected > 0 else 0.0
    phone_pct = (phone_sel / num_selected) * 100.0 if num_selected > 0 else 0.0
    
    max_sel_share_pct = max(email_pct, sms_pct, phone_pct)
    cap_respected = max_sel_share_pct <= 40.0 + 1e-4

    c17 = record_check("CHK_DST_01", "Channel Distribution", "Max channel share <= 40%", "<= 40.0%", f"{max_sel_share_pct:.2f}%", cap_respected)
    all_checks.append(c17)

    print("CHANNEL DISTRIBUTION")
    print("--------------------")
    print(f"Email: {email_sel} ({email_pct:.2f}%)")
    print(f"SMS: {sms_sel} ({sms_pct:.2f}%)")
    print(f"Phone Call: {phone_sel} ({phone_pct:.2f}%)")
    print(f"Maximum channel share: {max_sel_share_pct:.2f}%")
    print(f"40% cap respected: {'YES' if cap_respected else 'NO'}")
    dst_status = "PASS" if cap_respected else "FAIL"
    print(f"Status: {dst_status}")

    # =========================================================================
    # STEP 8 — EXISTING ADAPTER
    # =========================================================================
    print("\n--- STEP 8: RUNNING EXISTING ADAPTER LAYER ---")
    merge_keys = ["patient_id", "plan_id", "care_gap", "intervention_type"]
    best_full_df = pd.merge(
        best_cands_df,
        candidate_df.drop(columns=["probability_score", "calibrated_probability"], errors="ignore"),
        on=merge_keys,
        how="inner"
    )

    optimizer_input_df = ml_to_optimizer_input(
        rule_engine_df=best_full_df,
        ml_output_df=best_cands_df
    )
    opt_in_rows = len(optimizer_input_df)
    
    req_opt_cols = [
        "member_id", "plan_id", "member_name", "age", "gender",
        "care_gap", "intervention_type", "closure_probability",
        "uncertainty_lower", "uncertainty_upper", "measure_id",
        "measure_weight", "performance_opportunity", "eligible", "gap_open"
    ]
    missing_opt_cols = [c for c in req_opt_cols if c not in optimizer_input_df.columns]

    c18 = record_check("CHK_ADP_01", "Adapter", "Optimizer input rows == 50", 50, opt_in_rows, opt_in_rows == 50)
    c19 = record_check("CHK_ADP_02", "Adapter", "Required optimizer fields present", [], missing_opt_cols, len(missing_opt_cols) == 0)
    all_checks.extend([c18, c19])

    print("ADAPTER")
    print("-------")
    print(f"Input rows: {num_selected}")
    print(f"Optimizer input rows: {opt_in_rows}")
    print(f"Missing required optimizer fields: {missing_opt_cols if missing_opt_cols else 'None'}")
    adp_status = "PASS" if (opt_in_rows == 50 and len(missing_opt_cols) == 0) else "FAIL"
    print(f"Status: {adp_status}")

    # =========================================================================
    # STEP 9 — EXISTING OPTIMIZER
    # =========================================================================
    print("\n--- STEP 9: RUNNING EXISTING OPTIMIZER ---")
    opt_output_df = run_optimizer(
        df=optimizer_input_df,
        max_selected_members=50
    )
    final_members_count = len(opt_output_df)
    dup_members = final_members_count - opt_output_df["member_id"].nunique() if not opt_output_df.empty else 0
    
    # Calculate intervention overrides (Optimizer intervention vs ML selected intervention for member)
    opt_overrides = 0
    # Create expected intervention mapping from highest robust_quality_impact
    qi_df = optimizer_input_df.copy()
    perf_opp = pd.to_numeric(qi_df.get("performance_opportunity", pd.Series(1.0, index=qi_df.index)), errors="coerce").fillna(1.0)
    meas_wt = pd.to_numeric(qi_df.get("measure_weight", pd.Series(1.0, index=qi_df.index)), errors="coerce").fillna(1.0)
    unc_lower = pd.to_numeric(qi_df.get("uncertainty_lower", pd.Series(0.5, index=qi_df.index)), errors="coerce").fillna(0.5)
    qi_df["_quality_impact"] = perf_opp * meas_wt
    qi_df["_robust_quality_impact"] = qi_df["_quality_impact"] * unc_lower
    qi_sorted = qi_df.sort_values(by=["member_id", "_robust_quality_impact", "closure_probability"], ascending=[True, False, False])
    expected_interv_map = qi_sorted.drop_duplicates(subset=["member_id"]).set_index("member_id")["intervention_type"].to_dict()

    for _, r in opt_output_df.iterrows():
        m_id = r["member_id"]
        rec = r["recommended_intervention"]
        exp = expected_interv_map.get(m_id, rec)
        if rec != exp:
            opt_overrides += 1

    c20 = record_check("CHK_OPT_01", "Optimizer", "Final members <= 50", "<= 50", final_members_count, final_members_count <= 50)
    c21 = record_check("CHK_OPT_02", "Optimizer", "Duplicate members == 0", 0, dup_members, dup_members == 0)
    c22 = record_check("CHK_OPT_03", "Optimizer", "Intervention overrides == 0", 0, opt_overrides, opt_overrides == 0)
    all_checks.extend([c20, c21, c22])

    print("OPTIMIZER")
    print("---------")
    print(f"Optimizer input rows: {opt_in_rows}")
    print(f"Maximum allowed members: 50")
    print(f"Final selected members: {final_members_count}")
    print(f"Duplicate members: {dup_members}")
    print(f"Intervention overrides: {opt_overrides}")
    opt_status = "PASS" if (final_members_count <= 50 and dup_members == 0 and opt_overrides == 0) else "FAIL"
    print(f"Status: {opt_status}")

    # =========================================================================
    # STEP 10 — FINAL OUTPUT VALIDATION
    # =========================================================================
    print("\n--- STEP 10: VALIDATING FINAL OUTPUT ---")
    null_count = opt_output_df.isna().sum().sum() if not opt_output_df.empty else 0
    final_channels_valid = set(opt_output_df["recommended_intervention"].unique()).issubset(valid_channels) if not opt_output_df.empty else True
    
    c23 = record_check("CHK_OUT_01", "Final Output", "Zero null values in output", 0, null_count, null_count == 0)
    c24 = record_check("CHK_OUT_02", "Final Output", "All recommended interventions valid", True, bool(final_channels_valid), final_channels_valid)
    all_checks.extend([c23, c24])

    # =========================================================================
    # STEP 11 — FULL TEST REPORT EXCEL GENERATION
    # =========================================================================
    print("\n--- STEP 11: GENERATING EXCEL REPORT ---")
    os.makedirs(os.path.dirname(output_report_path), exist_ok=True)
    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    pass_font = Font(name="Calibri", size=11, bold=True, color="375623")
    fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fail_font = Font(name="Calibri", size=11, bold=True, color="C65911")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    bold_font = Font(name="Calibri", size=11, bold=True)

    passed_count = sum(1 for c in all_checks if c["Status"] == "PASS")
    failed_count = sum(1 for c in all_checks if c["Status"] == "FAIL")
    overall_status = "PASS" if failed_count == 0 else "FAIL"

    # Sheet 1: Summary
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.append(["CARE-GAP PIPELINE — UNSEEN 50-COMBINATION TEST REPORT"])
    ws_sum.cell(row=1, column=1).font = title_font
    ws_sum.append([])
    
    summary_items = [
        ["Test Dataset", "Unseen_50_Test_Data_New_Combinations (1).xlsx"],
        ["Original contexts", num_rows],
        ["Candidate rows", total_cands],
        ["ML prediction rows", preds_count],
        ["Calibrated rows", cal_count],
        ["Selected interventions", num_selected],
        ["Clear winners", clear_winners],
        ["Near ties", near_ties],
        ["Clear winners changed", clear_changed],
        ["Email", f"{email_sel} / {email_pct:.2f}%"],
        ["SMS", f"{sms_sel} / {sms_pct:.2f}%"],
        ["Phone Call", f"{phone_sel} / {phone_pct:.2f}%"],
        ["Optimizer input", opt_in_rows],
        ["Final selected members", final_members_count],
        ["Duplicate members", dup_members],
        ["Intervention overrides", opt_overrides],
        ["Validation checks passed", passed_count],
        ["Validation checks failed", failed_count],
        ["Overall Test Status", overall_status]
    ]

    for k, v in summary_items:
        r_idx = ws_sum.max_row + 1
        ws_sum.append([k, v])
        ws_sum.cell(row=r_idx, column=1).font = bold_font
        if k == "Overall Test Status":
            c_val = ws_sum.cell(row=r_idx, column=2)
            c_val.font = pass_font if v == "PASS" else fail_font
            c_val.fill = pass_fill if v == "PASS" else fail_fill

    def add_table_sheet(wb, title, df, format_headers=True):
        ws = wb.create_sheet(title=title)
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        if format_headers:
            for c in range(1, len(df.columns) + 1):
                cell = ws.cell(row=1, column=c)
                cell.fill = header_fill
                cell.font = header_font
        return ws

    # Sheet 2: Input_Validation
    add_table_sheet(wb, "Input_Validation", raw_input_df)

    # Sheet 3: Candidate_Validation
    add_table_sheet(wb, "Candidate_Validation", candidate_df)

    # Sheet 4: ML_Results
    add_table_sheet(wb, "ML_Results", ml_output_df)

    # Sheet 5: Calibration
    add_table_sheet(wb, "Calibration", ml_output_df[["patient_id", "care_gap", "intervention_type", "probability_score", "calibrated_probability"]])

    # Sheet 6: Intervention_Selection
    add_table_sheet(wb, "Intervention_Selection", audit_df)

    # Sheet 7: Channel_Distribution
    dist_df = pd.DataFrame([
        {"Channel": "Email", "Count": email_sel, "Percentage": f"{email_pct:.2f}%", "Max Cap": "40%", "Cap Respected": "YES" if email_pct <= 40.0 else "NO"},
        {"Channel": "SMS", "Count": sms_sel, "Percentage": f"{sms_pct:.2f}%", "Max Cap": "40%", "Cap Respected": "YES" if sms_pct <= 40.0 else "NO"},
        {"Channel": "Phone Call", "Count": phone_sel, "Percentage": f"{phone_pct:.2f}%", "Max Cap": "40%", "Cap Respected": "YES" if phone_pct <= 40.0 else "NO"},
    ])
    add_table_sheet(wb, "Channel_Distribution", dist_df)

    # Sheet 8: Optimizer_Validation
    add_table_sheet(wb, "Optimizer_Validation", opt_output_df)

    # Sheet 9: Final_Output
    add_table_sheet(wb, "Final_Output", opt_output_df)

    # Sheet 10: Failed_Checks
    failed_df = pd.DataFrame(failed_checks_list) if failed_checks_list else pd.DataFrame(columns=["Check ID", "Stage", "Description", "Expected", "Actual", "Status", "Remarks"])
    add_table_sheet(wb, "Failed_Checks", failed_df)

    wb.save(output_report_path)
    print(f"Report successfully saved: {output_report_path}")

    # =========================================================================
    # STEP 12 — TERMINAL REPORT
    # =========================================================================
    print("\n" + "=" * 60)
    print("UNSEEN NEW-COMBINATION TEST — FINAL RESULT")
    print("=" * 60)
    print(f"Dataset:                     Unseen_50_Test_Data_New_Combinations (1).xlsx")
    print(f"Input contexts:              {num_rows}")
    print(f"Candidates:                  {total_cands}")
    print(f"ML predictions:              {preds_count}")
    print(f"Calibrated predictions:      {cal_count}")
    print(f"Selected interventions:      {num_selected}")
    print(f"Clear winners:               {clear_winners}")
    print(f"Near ties:                   {near_ties}")
    print(f"Clear winners changed:       {clear_changed}")
    print(f"Email:                       {email_sel} ({email_pct:.2f}%)")
    print(f"SMS:                         {sms_sel} ({sms_pct:.2f}%)")
    print(f"Phone Call:                  {phone_sel} ({phone_pct:.2f}%)")
    print(f"Optimizer input:             {opt_in_rows}")
    print(f"Final members:               {final_members_count}")
    print(f"Duplicate members:           {dup_members}")
    print(f"Intervention overrides:      {opt_overrides}")
    print(f"Validation checks passed:    {passed_count}")
    print(f"Validation checks failed:    {failed_count}")
    print(f"OVERALL STATUS:              {overall_status}")
    print(f"Report:                      {output_report_path}")
    print("=" * 60)

    return overall_status


if __name__ == "__main__":
    run_unseen_50_test()
